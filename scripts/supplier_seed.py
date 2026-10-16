"""Generar INSERT privado desde las columnas de compras, sin conectar a la base.

python scripts/supplier_seed.py libro.xlsx --output artifacts/proveedores.sql
Las equivalencias específicas del negocio se pasan por --aliases archivo.json.
"""
import argparse
import csv
import json
from pathlib import Path
from openpyxl import load_workbook


def name_key(value):
    return " ".join(value.casefold().split())


def sql_literal(value):
    if "\x00" in value:
        raise ValueError("Un nombre contiene un carácter inválido")
    return "'" + value.replace("'", "''") + "'"


def suppliers_from_excel(path, aliases=None):
    aliases = {name_key(key): value.strip() for key, value in (aliases or {}).items()}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["ingresos y egresos"]
        sources = []
        names = {}
        for row in sheet.iter_rows(min_row=2, max_col=6):
            for index in (2, 4):
                cell = row[index]
                if not isinstance(cell.value, str) or not cell.value.strip():
                    continue
                original = cell.value
                normalized = aliases.get(name_key(original), " ".join(original.split()).title())
                if not normalized or len(normalized) > 150:
                    raise ValueError(f"Revisar nombre en {cell.coordinate}")
                key = name_key(normalized)
                names.setdefault(key, normalized)
                sources.append((cell.coordinate, original, normalized))
        return sorted(names.items()), sources
    finally:
        workbook.close()


def build_sql(names):
    rows = [f"  ({sql_literal(name)}, {sql_literal(key)}, '', '', true, 1)" for key, name in names]
    if not rows:
        raise ValueError("No se encontraron nombres en las columnas de compras")
    return """-- Proveedores normalizados. Revisables desde Gestión > Proveedores.
-- Ejecutar en la base del negocio con suppliers creado por Alembic.
-- No modifica proveedores existentes. Se puede ejecutar más de una vez.
BEGIN;
SET LOCAL standard_conforming_strings = on;
INSERT INTO suppliers (name, name_key, contact, notes, active, version)
VALUES
""" + ",\n".join(rows) + "\nON CONFLICT (name_key) DO NOTHING;\nCOMMIT;\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--aliases", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    aliases = json.loads(args.aliases.read_text(encoding="utf-8")) if args.aliases else {}
    names, sources = suppliers_from_excel(args.workbook, aliases)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(build_sql(names), encoding="utf-8")
    mapping = args.output.with_suffix(".csv")
    with mapping.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, delimiter=";")
        writer.writerow(("Celda", "Nombre original", "Proveedor normalizado"))
        writer.writerows(sources)
    print(f"{len(names)} proveedores; {len(sources)} referencias; INSERT y equivalencias generados.")


if __name__ == "__main__":
    main()
