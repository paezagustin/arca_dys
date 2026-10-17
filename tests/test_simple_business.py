"""Reglas de la simplificación: acceso local, caja diaria y exportación conciliable."""
from datetime import date
from decimal import Decimal
from io import BytesIO
import re
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, func, text, inspect
from sqlalchemy.orm import Session
from app.main import app
from app.core.database import engine
from app.models.business import BusinessEntry, Supplier
from app.models.invoice import Invoice
from app.models.settings import BillingPreview
from app.services.billing import BillingService
from app.services.business import save_entry, save_supplier
from app.services.listings import business_today
from tests.test_business import values
from tests.test_migrations import scratch_database, migrate


def csrf(page):
    return re.search(r'name="csrf_token" value="([^"]+)"', page.text).group(1)


def test_guest_can_save_sale_edit_and_export_without_login():
    with TestClient(app) as guest:
        page = guest.get("/business/entries/new?kind=sale")
        assert page.status_code == 200
        assert f'value="{business_today()}"' in page.text
        for field in ("supplier_id", "description", "expense_type", "category", "medium", "classification", "reason", "active"):
            assert f'name="{field}"' not in page.text
        assert "Tipo de registro" not in page.text and "Otro ingreso" not in page.text
        data = values(csrf_token=csrf(page))
        assert guest.post("/business/entries/new", data=data).status_code == 200
        with Session(engine) as db:
            identifier = db.scalar(select(BusinessEntry.id))
        data.update(version="1", cash_amount="3200.23")
        assert guest.post(f"/business/entries/{identifier}", data=data).status_code == 200
        assert guest.get("/dashboard?year=2026").status_code == 200
        exported = guest.get("/dashboard/export?year=2026&month=1")
        assert exported.status_code == 200
        sheet = load_workbook(BytesIO(exported.content), data_only=True).active
        assert sheet["B11"].value == 3200.23
        nav = re.search(r'<nav aria-label="Principal">(.*?)</nav>', page.text, re.S).group(1)
        for label in ("Facturación", "Gestión", "Resumen"):
            assert label in nav
        for label in ("Cobros MP", "Auditoría", "Configuración", "Conexión MP", "Pruebas"):
            assert label not in nav
        for url in ("/page/payments", "/page/reports-stadistics", "/settings", "/settings/mp", "/dev"):
            assert guest.get(url, follow_redirects=False).status_code == 303


def test_anonymous_preview_is_private_to_browser_and_confirmation_is_idempotent():
    with TestClient(app) as guest, TestClient(app) as other:
        token = csrf(guest.get("/page/invoice/manual"))
        created = guest.post("/ARCA/invoice/manual/preview", data={"amount": "10456.78", "csrf_token": token}, follow_redirects=False)
        assert created.status_code == 303
        url = created.headers["location"]
        assert other.get(url).status_code == 404
        other_token = csrf(other.get("/page/invoice/manual"))
        assert other.post(url + "/confirm", data={"csrf_token": other_token}).status_code == 422
        assert guest.post(url + "/confirm", data={}).status_code == 403
        for _ in range(2):
            assert guest.post(url + "/confirm", data={"csrf_token": token}).status_code == 200
        assert BillingService().process_one_queued()
        with Session(engine) as db:
            invoice = db.scalar(select(Invoice))
            assert db.scalar(select(func.count()).select_from(Invoice)) == 1
            assert invoice.total_amount == Decimal("10456.78")
            assert guest.get(f"/ARCA/invoice/{invoice.id}/pdf").content.startswith(b"%PDF")
        history = guest.get("/page/billing-requests")
        assert history.status_code == 200 and "Historial de solicitudes" in history.text


def test_expense_type_is_required_and_search_selects_by_id_with_literal_wildcards(client):
    with Session(engine) as db:
        supplier = save_supplier(db, {"name": "Ficticio O'Brien 100%"}, "test")
        supplier_id = supplier.id
        save_supplier(db, {"name": "Ficticio distinto"}, "test")
    found = client.get("/business/suppliers/search", params={"q": "100%"}).json()
    assert found == [{"id": supplier_id, "name": "Ficticio O'Brien 100%"}]
    assert client.post("/business/entries/new", data=values(kind="expense", expense_type="invalid")).status_code == 422
    assert client.post("/business/entries/new", data=values(kind="expense", expense_type="black", supplier_id=str(supplier_id))).status_code == 200
    with Session(engine) as db:
        assert db.scalar(select(BusinessEntry.expense_type)) == "black"
    assert client.post(f"/business/suppliers/{supplier_id}/delete", data={"version": "1"}).status_code == 422
    with Session(engine) as db:
        unused = db.scalar(select(Supplier).where(Supplier.id != supplier_id))
        unused_id = unused.id
    assert client.post(f"/business/suppliers/{unused_id}/delete", data={"version": "1"}).status_code == 200
    assert client.get(f"/business/suppliers/{unused_id}").status_code == 404


def test_export_has_six_columns_and_reconciles_multiple_expenses_dates_zero_and_text(client):
    with Session(engine) as db:
        supplier = save_supplier(db, {"name": "=SUM(1,1)"}, "test")
        for data in [values(entry_date="2026-01-01", amount="1234.56"), values(entry_date="2026-01-02", amount="0"),
            values(entry_date="2026-01-01", kind="expense", amount="100.01", supplier_id=str(supplier.id)),
            values(entry_date="2026-01-01", kind="expense", amount="200.02", description="Texto"),
            values(entry_date="2026-01-01", kind="expense", amount="300.03", expense_type="black"),
            values(entry_date="2026-02-01", kind="expense", amount="999.99")]:
            save_entry(db, data, "test")
    response = client.get("/dashboard/export?year=2026&month=1")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content), data_only=False).active
    assert sheet.title == "ingresos y egresos" and sheet.max_column == 6 and sheet.max_row == 33
    assert [cell.value for cell in sheet[1]] == ["FECHA", "VENTAS DEL DIA", "COMPRAS EN BLANCO", "MONTO", "COMPRAS EN NEGRO", "MONTO2"]
    assert sheet["A2"].value.date() == date(2026, 1, 1)
    assert sheet["B2"].value == 1234.56 and sheet["B3"].value is None and sheet["B4"].value == 0
    assert sheet["C2"].data_type == "s" and sheet["C2"].value == "=SUM(1,1)"
    assert sheet["B2"].number_format == '"$ "#,##0.00'
    assert sheet["A1"].fill.fgColor.rgb == "00ADC5E7"
    totals = {col: sum((Decimal(str(row[col-1].value)) for row in sheet.iter_rows(min_row=2) if row[col-1].value is not None), Decimal(0)) for col in (2, 4, 6)}
    assert totals == {2: Decimal("1234.56"), 4: Decimal("300.03"), 6: Decimal("300.03")}
    page = client.get("/dashboard?year=2026&month=1").text
    for label in ("Ventas", "Egresos en blanco", "Egresos en negro", "Total Egresos", "Ganancia"):
        assert label in page
    assert "$ 634,50" in page and "Otros ingresos" not in page


def test_export_rejects_unknown_legacy_expense_and_excel_precision_loss(client):
    with Session(engine) as db:
        row = save_entry(db, values(kind="expense"), "test")
        row.expense_type = None
        db.commit()
    assert client.get("/dashboard/export?year=2026&month=1").status_code == 422
    with Session(engine) as db:
        row = db.scalar(select(BusinessEntry))
        row.expense_type, row.amount = "white", Decimal("100000000000000.01")
        db.commit()
    response = client.get("/dashboard/export?year=2026&month=1")
    assert response.status_code == 422 and "precisión" in response.text


def test_simple_migration_preserves_amounts_and_maps_only_known_types(scratch_database):
    migrate(scratch_database, "0009_reports")
    with scratch_database.begin() as db:
        db.execute(text("""INSERT INTO business_entries(entry_key,kind,entry_date,amount,category,medium,description,classification,origin,active,version,created_at)
        VALUES ('test-white','expense','2026-01-01',12.34,'Compras','cash','','COMPRAS EN BLANCO','manual',true,1,now()),
        ('test-unknown','expense','2026-01-02',56.78,'Luz','cash','Local','','manual',true,1,now())"""))
    migrate(scratch_database)
    with scratch_database.connect() as db:
        assert db.execute(text("SELECT sum(amount) FROM business_entries")).scalar_one() == Decimal("69.12")
        assert db.execute(text("SELECT expense_type FROM business_entries WHERE entry_key='test-white'")).scalar_one() == "white"
        assert db.execute(text("SELECT expense_type,description FROM business_entries WHERE entry_key='test-unknown'")).one() == (None, "Local · Luz")
        assert "business_changes" not in inspect(db).get_table_names()
        assert "settings_changes" not in inspect(db).get_table_names()
        assert not {"medium", "category", "classification"} & {col["name"] for col in inspect(db).get_columns("business_entries")}


def test_simple_migration_does_not_delete_existing_other_income(scratch_database):
    migrate(scratch_database, "0009_reports")
    with scratch_database.begin() as db:
        db.execute(text("""INSERT INTO business_entries(entry_key,kind,entry_date,amount,category,medium,description,classification,origin,active,version,created_at)
        VALUES ('test-income','income','2026-01-01',12.34,'Reintegro','cash','','','manual',true,1,now())"""))
    with pytest.raises(RuntimeError, match="otros ingresos"):
        migrate(scratch_database)
    with scratch_database.connect() as db:
        assert db.execute(text("SELECT amount FROM business_entries")).scalar_one() == Decimal("12.34")
        assert db.execute(text("SELECT version_num FROM alembic_version")).scalar_one() == "0009_reports"
