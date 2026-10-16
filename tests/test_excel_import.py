"""Libros sintéticos en memoria; nunca se incluye el Excel privado en la suite."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
import pytest
from openpyxl import Workbook
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from app.core.database import engine
from app.models.imports import ImportBatch, ImportRow
from app.models.business import BusinessEntry, Supplier
from app.models.payment import Payment
from app.models.invoice import Invoice
from app.services.excel_reader import read_excel, HEADERS
from app.services.excel_import import prepare_import, confirm_import, resolve_row, assign_supplier
from app.services.business import save_supplier
from app.services.business import save_entry
from tests.test_business import values


def workbook_bytes(changed=False, extra=False):
    # Se genera únicamente un fixture técnico descartable, no un archivo para el usuario.
    book = Workbook()
    detail = book.active
    detail.title = "ingresos y egresos"
    detail.append(HEADERS)
    detail.append([datetime(2026, 1, 1), 101 if changed else 100, "Proveedor Uno", 25, "Proveedor Dos", 75])
    detail.append([datetime(2026, 1, 2), 0, "Importe ambiguo", ",", None, None])
    detail.append([datetime(2026, 1, 3), None, None, None, "Sin importe", None])
    detail.append([datetime(2026, 1, 4), 10, None, 20, None, None])
    if extra:
        detail.append([datetime(2026, 1, 5), 20])
    dash = book.create_sheet("dashboard")
    for key, value in {"C2": "MESES", "D2": "VENTAS", "E2": "COMPRAS EN BLANCO", "F2": "COMPRAS EN NEGRO",
        "C3": datetime(2026, 1, 1), "D3": "=SUM('ingresos y egresos'!B2:B5)",
        "C65": "AFIP-FACTURACIÓN", "C66": "MESES", "D66": "RECIBIDO", "E66": "EMITIDO",
        "C67": datetime(2022, 1, 1), "D67": 123, "E67": 456, "F67": "=E67-D67",
        "H67": datetime(2022, 1, 1), "I67": 124, "J67": 456}.items():
        dash[key] = value
    excluded = book.create_sheet("sueldo privado")
    excluded["A1"] = "NO IMPORTAR ESTE DATO"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def prepare(db, content=None):
    content = content or workbook_bytes()
    return prepare_import(db, content, "Libro-ficticio-2026", "ejemplo.xlsx", "test", read_excel(content, "Libro-ficticio-2026"))


def test_reader_preserves_zero_originals_classification_and_distinct_manual_blocks():
    rows, report = read_excel(workbook_bytes(), "Libro-ficticio-2026")
    candidates = [row for row in rows if row["kind"] in {"sale", "expense"}]
    assert sum(row["state"] == "review" for row in candidates) == 2
    assert any(row["proposed"].get("amount") == "0.00" for row in candidates)
    assert report["source_totals"] == {"2026-01": {"sales": "110.00", "D": "45.00", "F": "75.00"}}
    historical = next(row for row in rows if row["kind"] == "historical" and row["row_number"] == 67)
    assert historical["raw"]["D67"]["value"] == 123
    assert historical["raw"]["I67"]["value"] == 124
    assert historical["raw"]["F67"]["value"] == "=E67-D67"
    assert all("NO IMPORTAR" not in str(row) for row in rows)
    assert report["dashboard_controls"][0]["difference"] is None


def test_preview_confirm_and_reupload_are_independent_and_idempotent(client):
    data = workbook_bytes()
    response = client.post("/business/imports", data={"dataset": "Libro-ficticio-2026"},
        files={"workbook": ("ejemplo.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200 and "Revisar importación" in response.text
    identifier = str(response.url).split("/")[-1]
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(BusinessEntry)) == 0
    for _ in range(2):
        assert client.post(f"/business/imports/{identifier}/confirm").status_code == 200
    again = client.post("/business/imports", data={"dataset": "Libro-ficticio-2026"}, files={"workbook": ("otro-nombre.xlsx", data)})
    assert str(again.url).endswith(identifier)
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(ImportBatch)) == 1
        assert db.scalar(select(func.count()).select_from(BusinessEntry)) == 6
        assert db.scalar(select(func.sum(BusinessEntry.amount)).where(BusinessEntry.kind == "expense")) == Decimal(120)
        assert db.scalar(select(func.count()).select_from(Payment)) == 0
        assert db.scalar(select(func.count()).select_from(Invoice)) == 0
    assert "Revisar fila" in client.get(f"/business/imports/{identifier}?state=review").text


def test_review_keeps_raw_and_audits_decision_then_imports_only_after_confirmation(client):
    with Session(engine) as db:
        batch = prepare(db)
        identifier = batch.id
        confirm_import(db, identifier, "test")
        row = db.scalar(select(ImportRow).where(ImportRow.batch_id == identifier, ImportRow.slot == "D", ImportRow.row_number == 3))
        row_id, raw = row.id, row.raw
    response = client.post(f"/business/imports/{identifier}/rows/{row_id}", data={"entry_date": "2026-01-02", "amount": "25",
        "description": "Compra confirmada", "category": "Compras", "reason": "Revisado contra original", "action": "ready"})
    assert response.status_code == 200
    with Session(engine) as db:
        row = db.get(ImportRow, row_id)
        assert row.raw == raw and row.state == "ready"
        assert row.resolution["history"][0]["actor"] == "admin"
        assert db.scalar(select(func.count()).select_from(BusinessEntry)) == 6
        confirm_import(db, identifier, "test")
        assert db.scalar(select(func.sum(BusinessEntry.amount)).where(BusinessEntry.kind == "expense")) == Decimal(145)
    assert client.post(f"/business/imports/{identifier}/rows/{row_id}", data={"reason": "No alterar importada", "action": "skip"}).status_code == 422


def test_omit_requires_reason_and_does_not_turn_missing_amount_into_zero(client):
    with Session(engine) as db:
        batch = prepare(db)
        identifier = batch.id
        row_id = db.scalar(select(ImportRow.id).where(ImportRow.batch_id == identifier, ImportRow.slot == "F", ImportRow.row_number == 4))
    assert client.post(f"/business/imports/{identifier}/rows/{row_id}", data={"action": "skip"}).status_code == 422
    assert client.post(f"/business/imports/{identifier}/rows/{row_id}", data={"action": "skip", "reason": "Sin comprobante"}).status_code == 200
    with Session(engine) as db:
        confirm_import(db, identifier, "test")
        row = db.get(ImportRow, row_id)
        assert row.state == "skipped" and row.entry_id is None


def test_new_file_preserves_same_origins_and_flags_changed_values():
    with Session(engine) as db:
        first = prepare(db)
        confirm_import(db, first.id, "test")
        second = prepare(db, workbook_bytes(changed=True, extra=True))
        confirm_import(db, second.id, "test")
        assert db.scalar(select(func.count()).select_from(BusinessEntry)) == 7
        row = db.scalar(select(ImportRow).where(ImportRow.batch_id == second.id, ImportRow.slot == "B", ImportRow.row_number == 2))
        assert row.state == "review" and row.resolution["existing_entry_id"]
        assert row.proposed["amount"] == "101.00"
        assert db.get(BusinessEntry, row.resolution["existing_entry_id"]).amount == Decimal(100)
        assert db.scalar(select(func.count()).select_from(ImportRow).where(ImportRow.batch_id == second.id, ImportRow.state == "already")) == 5


def test_concurrent_confirm_does_not_duplicate():
    with Session(engine) as db:
        identifier = prepare(db).id
    def confirm(_):
        with Session(engine) as db:
            return confirm_import(db, identifier, "test").id
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert list(pool.map(confirm, range(2))) == [identifier, identifier]
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(BusinessEntry)) == 6


def test_suppliers_are_suggestions_until_explicit_mapping(client):
    with Session(engine) as db:
        batch = prepare(db)
        supplier = save_supplier(db, {"name": "Proveedor Uno"}, "test")
        identifier, supplier_id = batch.id, supplier.id
    page = client.get(f"/business/imports/{identifier}/suppliers")
    assert "Nombres parecidos: Proveedor Uno" in page.text
    assert client.post(f"/business/imports/{identifier}/suppliers", data={"label": "Proveedor Uno", "supplier_id": str(supplier_id)}).status_code == 200
    with Session(engine) as db:
        confirm_import(db, identifier, "test")
        linked = list(db.scalars(select(BusinessEntry).where(BusinessEntry.supplier_id.isnot(None))))
        assert len(linked) == 1 and linked[0].description == "Proveedor Uno"
        assert db.scalar(select(func.count()).select_from(Supplier)) == 1


@pytest.mark.parametrize("content", [b"", b"not a workbook", b"x" * (5 * 1024 * 1024 + 1)])
def test_malformed_and_oversized_files_are_rejected(content):
    with pytest.raises(ValueError):
        read_excel(content, "Prueba")


def test_original_lists_are_paginated(client):
    with Session(engine) as db:
        identifier = prepare(db).id
    assert client.get(f"/business/imports/{identifier}?state=source&page=2").status_code == 200
    assert client.get(f"/business/imports/{identifier}?state=unknown").status_code == 422
    history = client.get(f"/dashboard/history?batch={identifier}")
    assert history.status_code == 200
    assert "$ 123,00" in history.text and "$ 124,00" in history.text
    assert "$ 333,00" in history.text and "$ 332,00" in history.text
    assert "1 registros" in history.text
    assert "Izquierdo" in history.text and "Derecho" in history.text


def test_import_does_not_replace_a_manual_sale():
    with Session(engine) as db:
        manual = save_entry(db, values(entry_date="2026-01-01", amount="999"), "test")
        identifier = prepare(db).id
        confirm_import(db, identifier, "test")
        row = db.scalar(select(ImportRow).where(ImportRow.batch_id == identifier, ImportRow.slot == "B", ImportRow.row_number == 2))
        assert row.state == "review" and row.entry_id is None
        assert db.get(BusinessEntry, manual.id).amount == Decimal(999)


def test_xml_entities_are_not_expanded():
    from zipfile import ZipFile, ZIP_DEFLATED
    source, output = BytesIO(workbook_bytes()), BytesIO()
    with ZipFile(source) as original, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "xl/workbook.xml":
                content = b'<!DOCTYPE workbook [<!ENTITY external SYSTEM "file:///etc/passwd">]>' + content
            target.writestr(item, content)
    with pytest.raises(ValueError):
        read_excel(output.getvalue(), "Prueba")
