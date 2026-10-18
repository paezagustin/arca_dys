"""Correcciones de Gestión sobre PostgreSQL aislado, sin proveedores externos."""
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from app.main import app
from app.core.database import engine
from app.models.business import BusinessEntry
from app.models.data_import import DataImport, LegacyImportRow
from app.models.invoice import Invoice
from app.models.payment import Payment
from app.models.payment_invoice import PaymentInvoice
from app.services.business import delete_entry, save_entry, save_supplier
from app.services.metrics import period_metrics
from tests.test_business import values
from tests.test_simple_business import csrf


def test_guest_delete_sunday_keeps_original_and_allows_corrected_closing():
    data = values(entry_date="2026-01-11", cash_amount="100.01", mp_amount="20.02")
    with Session(engine) as db:
        identifier = save_entry(db, data, "test").id
    with TestClient(app) as guest:
        url = f"/business/entries/{identifier}"
        token = csrf(guest.get(url))
        assert guest.get(url + "/delete").status_code == 405
        assert guest.post(url + "/delete", data={"version": "1"}).status_code == 403
        for _ in range(2):
            result = guest.post(url + "/delete", data={"version": "1", "csrf_token": token})
            assert result.status_code == 200 and "Registro eliminado" in result.text
        assert guest.get(url).status_code == 404
        assert guest.post("/business/entries/new", data=dict(data, csrf_token=token)).status_code == 422
        assert guest.post(url, data=dict(data, version="2", csrf_token=token)).status_code == 422
        assert guest.post("/business/entries/new", data=dict(data, entry_key=str(uuid4()), csrf_token=token)).status_code == 200
    with Session(engine) as db:
        original = db.get(BusinessEntry, identifier)
        assert not original.active and original.version == 2
        assert original.entry_key == data["entry_key"] and original.entry_date == date(2026, 1, 11)
        assert (original.cash_amount, original.mp_amount, original.amount) == (Decimal("100.01"), Decimal("20.02"), Decimal("120.03"))
        assert len(list(db.scalars(select(BusinessEntry)))) == 2
        assert period_metrics(db, date(2026, 1, 1), date(2026, 1, 31))["sale"] == Decimal("120.03")


def test_delete_expense_preserves_import_and_fiscal_evidence_and_updates_reports(client):
    with Session(engine) as db:
        row = save_entry(db, values(kind="expense", amount="25.02", description="Original ficticio"), "test")
        row.origin = "excel"
        identifier = row.id
        batch_id = str(uuid4())
        db.add(DataImport(id=batch_id, dataset="synthetic", file_hash="a"*64,
                          filename="synthetic.json", state="applied", report={"test": True}, actor="test"))
        db.flush()
        db.add(LegacyImportRow(batch_id=batch_id, sheet="test", row_number=1, slot="expense", source_key="s",
                              fingerprint="f", kind="expense", raw={"original": "25.02"}, proposed={}, state="applied", entry_id=identifier))
        payment = Payment(id_payment_mp="SIM-DELETE-SAFE", mont=Decimal("500"), registered_arca=True,
                          status="approved", raw_json={"synthetic": "original"})
        invoice = Invoice(invoice_number="SIM-DELETE-SAFE", total_amount=Decimal("500"), status="success",
                          cae="SIMULADO", raw_afip_response={"synthetic": True})
        db.add_all([payment, invoice])
        db.flush()
        db.add(PaymentInvoice(payment_id=payment.id, invoice_id=invoice.id))
        db.commit()
        tables = ("payments", "invoices", "payment_invoice", "mp_payment_identities", "import_batches", "import_rows")
        def evidence():
            return {name: db.execute(text(f'SELECT row_to_json(t) FROM "{name}" t ORDER BY 1::text')).scalars().all() for name in tables}
        before = evidence()
        original = db.execute(text("SELECT row_to_json(t) FROM business_entries t WHERE id=:id"), {"id": identifier}).scalar_one()
    assert client.post(f"/business/entries/{identifier}/delete?kind=expense&supplier_id=none&page=2&redirect=https://example.com",
                       data={"version": "1"}, follow_redirects=False).headers["location"] == "/business?kind=expense&supplier_id=none&page=2&deleted=1"
    with Session(engine) as db:
        assert evidence() == before
        after = db.execute(text("SELECT row_to_json(t) FROM business_entries t WHERE id=:id"), {"id": identifier}).scalar_one()
        assert after == dict(original, active=False, version=2)
        summary = period_metrics(db, date(2026, 1, 1), date(2026, 1, 31))
        assert summary["expense"] is None
    assert client.get("/business?kind=expense").context["entries"].total == 0
    book = load_workbook(BytesIO(client.get("/dashboard/export?year=2026&month=1").content), data_only=True)
    assert all(row[3].value is None and row[5].value is None for row in book.active.iter_rows(min_row=2))


def test_delete_rejects_missing_and_stale_version_after_edit(client):
    with Session(engine) as db:
        identifier = save_entry(db, values(kind="expense"), "test").id
        save_entry(db, values(kind="expense", amount="250", version="1"), "test", identifier)
    for version in (None, "bad", "1"):
        assert client.post(f"/business/entries/{identifier}/delete", data={"version": version} if version else {}).status_code == 422
    assert client.post("/business/entries/999999/delete", data={"version": "1"}).status_code == 404
    with Session(engine) as db:
        assert db.get(BusinessEntry, identifier).active
        assert db.get(BusinessEntry, identifier).amount == Decimal("250")
    assert client.post(f"/business/entries/{identifier}/delete", data={"version": "2"}).status_code == 200


def test_concurrent_edit_delete_cannot_lose_correction():
    with Session(engine) as db:
        identifier = save_entry(db, values(kind="expense"), "test").id
    def change(action):
        with Session(engine) as db:
            try:
                if action == "delete":
                    delete_entry(db, identifier, "1")
                else:
                    save_entry(db, values(kind="expense", amount="250", version="1"), "test", identifier)
                return action
            except ValueError:
                db.rollback()
                return "conflict"
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(change, ("edit", "delete")))
    assert results.count("conflict") == 1
    with Session(engine) as db:
        row = db.get(BusinessEntry, identifier)
        assert row.version == 2
        assert (row.active, row.amount) == ((True, Decimal("250")) if "edit" in results else (False, Decimal("1500.25")))


def test_supplier_filter_totals_pagination_dates_archived_and_without_supplier(client):
    with Session(engine) as db:
        selected = save_supplier(db, {"name": "Ficticio A"}, "test")
        selected_id = selected.id
        other = save_supplier(db, {"name": "Ficticio B"}, "test")
        for n in range(27):
            save_entry(db, values(kind="expense", amount="10.01", supplier_id=str(selected_id)), "test")
        save_entry(db, values(kind="expense", amount="100", supplier_id=str(selected_id), entry_date="2025-12-31"), "test")
        save_entry(db, values(kind="expense", amount="777", supplier_id=str(other.id)), "test")
        save_entry(db, values(kind="expense", amount="25"), "test")
        save_entry(db, values(amount="1000"), "test")
        selected.active = False
        db.commit()
    query = f"supplier_id={selected_id}&kind=all&fecha_desde=2026-01-01&fecha_hasta=2026-01-31"
    page = client.get("/business?" + query)
    assert page.context["entries"].total == 27 and len(page.context["entries"].items) == 25
    assert page.context["totals"] == {"expense": Decimal("270.27")}
    assert "Ficticio A (archivado)" in page.text
    next_page = client.get("/business?" + query + "&page=2")
    assert len(next_page.context["entries"].items) == 2
    assert next_page.context["totals"] == page.context["totals"]
    assert f"supplier_id={selected_id}&amp;kind=all&amp;fecha_desde=2026-01-01&amp;fecha_hasta=2026-01-31&amp;page=2" in page.text
    assert client.get("/business?supplier_id=none").context["totals"] == {"expense": Decimal("25")}
    assert client.get(f"/business?kind=sale&supplier_id={selected_id}").context["entries"].total == 0
    assert client.get("/business").context["entries"].total == 31


@pytest.mark.parametrize("supplier", ["invalid", "-1", "0", "99999999999999999999999", "999999"])
def test_supplier_filter_rejects_invalid_selection(client, supplier):
    assert client.get("/business?supplier_id=" + supplier).status_code == 422
