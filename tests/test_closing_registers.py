from datetime import date, datetime, timedelta, time
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
import re
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, func, text
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError, DataError
from app.main import app
from app.core.database import engine
from app.models.business import BusinessEntry
from app.models.payment import Payment
from app.models.registers import SalaryPayment, CustomerCredit
from app.services.business import save_entry
from app.services.registers import save_register
from app.services.listings import business_today
from app.services.metrics import year_metrics, period_total, period_metrics
from tests.test_business import values
from tests.test_simple_business import csrf
from tests.test_migrations import scratch_database, migrate


def register(**changes):
    return dict({"entry_key": str(uuid4()), "entry_date": str(business_today()), "amount": "105.25",
        "employee": "Francisco", "customer": "Cliente ficticio", "observation": "Prueba"}, **changes)


def test_mp_suggestion_uses_only_local_approved_unbilled_payments_and_day_boundaries(client):
    today = business_today()
    with Session(engine) as db:
        for n, (status, billed, day, amount) in enumerate([
            ("approved", False, datetime.combine(today, time.min), "10.10"),
            ("approved", False, datetime.combine(today, time.max), "20.20"),
            ("approved", True, datetime.combine(today, time.min), "30"),
            ("refunded", False, datetime.combine(today, time.min), "40"),
            ("approved", False, datetime.combine(today-timedelta(days=1), time.max), "50"),
            ("approved", False, datetime.combine(today+timedelta(days=1), time.min), "60"),
            ("approved", False, None, "70"),
            ("approved", False, datetime.combine(today, time.min), "-80")]):
            db.add(Payment(id_payment_mp=f"SIM-CLOSING-{n}", status=status, registered_arca=billed,
                date_approved=day, mont=Decimal(amount)))
        db.commit()
    assert client.get(f"/business/sales/mp-total?date={today}").json() == {"total": "30.30"}
    page = client.get("/business/entries/new?kind=sale")
    assert 'value="30.30" placeholder="0,00"' in page.text and 'name="cash_amount"' in page.text
    assert client.get("/business/sales/mp-total?date=2026-02-30").status_code == 422


def test_close_can_override_mp_total_server_sums_and_edit_does_not_recalculate(client):
    data = values(cash_amount="100.01", mp_amount="200.02", amount="999999")
    assert client.post("/business/entries/new", data=data).status_code == 200
    with Session(engine) as db:
        row = db.scalar(select(BusinessEntry))
        identifier = row.id
        assert (row.cash_amount, row.mp_amount, row.amount) == (Decimal("100.01"), Decimal("200.02"), Decimal("300.03"))
        db.add(Payment(id_payment_mp="SIM-LATER", status="approved", registered_arca=True,
            mont=Decimal("1000"), date_approved=datetime(2026,1,10)))
        db.commit()
    page = client.get(f"/business/entries/{identifier}")
    assert 'value="200.02" placeholder="0,00"' in page.text
    assert client.post(f"/business/entries/{identifier}", data=dict(data, version="1", cash_amount="100.02")).status_code == 200
    with Session(engine) as db:
        assert db.get(BusinessEntry, identifier).amount == Decimal("300.04")


@pytest.mark.parametrize("cash,mp", [("", ""), ("-1", "0"), ("NaN", "0"), ("0", "Infinity"),
    ("1.001", "0"), ("9999999999999999.99", "0.01")])
def test_invalid_breakdown_is_rejected_without_saving(client, cash, mp):
    assert client.post("/business/entries/new", data=values(cash_amount=cash, mp_amount=mp, amount="total ignorado")).status_code == 422
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(BusinessEntry)) == 0


@pytest.mark.parametrize("cash,mp", [("12.34", ""), ("", "12.34"), ("0", ""), ("", "0")])
def test_partial_close_treats_blank_side_as_zero_and_keeps_explicit_zero(client, cash, mp):
    response = client.post("/business/entries/new", data=values(cash_amount=cash, mp_amount=mp))
    assert response.status_code == 200
    with Session(engine) as db:
        row = db.scalar(select(BusinessEntry))
        assert (row.cash_amount, row.mp_amount, row.amount) == (
            Decimal(cash or "0"), Decimal(mp or "0"), Decimal(cash or "0")+Decimal(mp or "0"))


def test_new_public_sections_have_only_requested_fields_and_csrf():
    with TestClient(app) as guest:
        page = guest.get("/salaries")
        assert page.status_code == 200 and "Francisco" in page.text and "Franisco" not in page.text
        assert '<option value="mercado pago" selected>' in page.text
        assert f'value="{business_today()}"' in page.text
        nav = re.search(r'<nav aria-label="Principal">(.*?)</nav>', page.text, re.S).group(1)
        assert "Sueldos" in nav and "Fiados" in nav and "Configuración" not in nav
        for url in ("/salaries", "/fiados"):
            assert guest.post(url, data=register()).status_code == 403
            form = guest.get(url)
            assert guest.post(url, data=register(csrf_token=csrf(form))).status_code == 200
            for field in ("reason", "active", "category", "classification"):
                assert f'name="{field}"' not in form.text
        assert guest.post("/fiados/1/paid", data={"version": "1"}).status_code == 403


def test_salary_default_month_employee_date_filter_and_pagination(client):
    today = business_today()
    start = today.replace(day=1)
    with Session(engine) as db:
        for n in range(27):
            save_register(db, SalaryPayment, register(employee="Juan Cruz", amount=str(n+1)))
        save_register(db, SalaryPayment, register(employee="Francisco", amount="500"))
        save_register(db, SalaryPayment, register(entry_date=str(start-timedelta(days=1)), amount="999"))
    page = client.get("/salaries?employee=Juan+Cruz&page=2")
    assert page.status_code == 200 and "Página 2 de 2" in page.text and "$ 378,00" in page.text
    assert "$ 999,00" not in page.text and "$ 500,00" not in page.text
    page = client.get(f"/salaries?employee=Francisco&fecha_desde={start-timedelta(days=1)}&fecha_hasta={today}")
    assert "$ 1.499,00" in page.text
    assert client.get("/salaries?employee=Otra").status_code == 422
    assert client.get("/salaries?fecha_desde=2026-05-01&fecha_hasta=2026-01-01").status_code == 422


@pytest.mark.parametrize("changes", [{"employee": "Franisco"}, {"medium": "tarjeta"}, {"amount": "-1"}, {"amount": "0"}, {"amount": "0.001"}, {"entry_date": "2026-02-30"}])
def test_invalid_salary_does_not_save(client, changes):
    assert client.post("/salaries", data=register(**changes)).status_code == 422
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(SalaryPayment)) == 0


def test_register_replays_and_stale_edits_are_safe(client):
    data = register()
    assert client.post("/salaries", data=data).status_code == 200
    assert client.post("/salaries", data=data).status_code == 200
    assert client.post("/salaries", data=dict(data, amount="999")).status_code == 422
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(SalaryPayment)) == 1
        identifier = db.scalar(select(SalaryPayment.id))
    update = dict(data, id=str(identifier), version="1", amount="200.25")
    assert client.post("/salaries", data=update).status_code == 200
    assert client.post("/salaries", data=update).status_code == 422
    with Session(engine) as db:
        assert db.get(SalaryPayment, identifier).amount == Decimal("200.25")


def test_credits_include_old_unpaid_allow_paid_edit_and_escape_text(client):
    data = register(entry_date="2020-01-01", customer="<script>Cliente</script>", observation="<img onerror=x>")
    assert client.post("/fiados", data=data).status_code == 200
    assert client.post("/fiados", data=data).status_code == 200
    page = client.get("/fiados")
    assert "&lt;script&gt;Cliente" in page.text and "&lt;img onerror=x&gt;" in page.text
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(CustomerCredit)) == 1
        row = db.scalar(select(CustomerCredit))
        identifier = row.id
        assert row.paid is False
        assert db.scalar(select(func.count()).select_from(BusinessEntry)) == 0
    assert client.post(f"/fiados/{identifier}/paid", data={"version": "99"}).status_code == 409
    for _ in range(2):
        assert client.post(f"/fiados/{identifier}/paid", data={"version": "1"}).status_code == 200
    assert "&lt;script&gt;Cliente" not in client.get("/fiados").text
    assert "&lt;script&gt;Cliente" in client.get("/fiados?paid=yes&customer=Cliente").text
    assert client.post("/fiados", data=dict(data, id=str(identifier), version="2", paid="0")).status_code == 200
    assert "&lt;script&gt;Cliente" in client.get("/fiados").text
    assert client.post("/fiados", data=register(customer="")).status_code == 422
    assert client.post("/fiados", data=register(paid="maybe")).status_code == 422


def test_metrics_and_export_reconcile_cash_mp_legacy_salaries_and_cross_year_range(client):
    with Session(engine) as db:
        save_entry(db, values(entry_date="2026-01-01", cash_amount="600.01", mp_amount="400.02"), "test")
        save_entry(db, values(kind="expense", amount="100.03", entry_date="2026-01-01"), "test")
        db.add(BusinessEntry(entry_key=str(uuid4()),kind="sale",entry_date=date(2026,1,2),amount=Decimal("200"),origin="excel",active=True,version=1,description=""))
        db.commit()
        save_register(db, SalaryPayment, register(entry_date="2026-01-01", amount="50.01"))
        save_register(db, SalaryPayment, register(entry_date="2026-01-02", employee="Juan Cruz", amount="25.02"))
        save_register(db, SalaryPayment, register(entry_date="2025-12-31", amount="10.01"))
        save_register(db, CustomerCredit, register(entry_date="2026-01-01", amount="99999"))
        january = year_metrics(db, 2026)[0]
        assert (january["sale_cash"],january["sale_mp"],january["sale_unknown"]) == (Decimal("600.01"),Decimal("400.02"),Decimal("200"))
        assert january["salary"] == Decimal("75.03") and january["result"] == Decimal("1024.97")
        assert january["employees"] == {"Francisco": Decimal("50.01"), "Juan Cruz": Decimal("25.02")}
        assert period_total(year_metrics(db, 2026))["result"] == january["result"]
        day = period_metrics(db, date(2026,1,1), date(2026,1,1))
        assert day["result"] == Decimal("849.99") and day["days"] == 1
        salary_only = period_metrics(db, date(2025,12,31), date(2025,12,31))
        assert salary_only["result"] == Decimal("-10.01")
    query = "?year=2026&fecha_desde=2025-12-31&fecha_hasta=2026-01-01"
    page = client.get("/dashboard"+query)
    assert page.status_code == 200 and "$ 839,98" in page.text and "Francisco" in page.text
    assert 'data-metric="sale_cash"' in page.text and "Sin desglose" not in page.text
    result = client.get("/dashboard/export"+query)
    assert result.status_code == 200
    book = load_workbook(BytesIO(result.content), data_only=True)
    assert book.sheetnames == ["ingresos y egresos", "Ventas", "Sueldos"]
    assert book.active.max_column == 6
    assert [c.value for c in book["Ventas"][2]][1:] == [600.01,400.02,1000.03]
    assert [r[2].value for r in book["Sueldos"].iter_rows(min_row=2)] == [10.01,50.01]
    assert client.get("/dashboard?fecha_desde=2026-01-01").status_code == 422
    assert client.get("/dashboard/export?fecha_desde=2026-02-01&fecha_hasta=2026-01-01").status_code == 422


def test_migration_preserves_legacy_sale_and_enforces_native_enums_and_breakdown(scratch_database):
    migrate(scratch_database, "0010_simple_business")
    with scratch_database.begin() as db:
        db.execute(text("""INSERT INTO business_entries(entry_key,kind,entry_date,amount,description,origin,active,version,created_at)
            VALUES ('SIM-OLD','sale','2026-01-01',123.45,'','manual',true,1,now())"""))
    migrate(scratch_database)
    with scratch_database.connect() as db:
        assert db.execute(text("SELECT amount,cash_amount,mp_amount FROM business_entries")).one() == (Decimal("123.45"),None,None)
        assert db.execute(text("SELECT enum_range(NULL::salary_employee)::text")).scalar_one() == '{Francisco,"Juan Cruz"}'
        assert db.execute(text("SELECT enum_range(NULL::salary_medium)::text")).scalar_one() == '{efectivo,"mercado pago"}'
    with pytest.raises(IntegrityError), scratch_database.begin() as db:
        db.execute(text("UPDATE business_entries SET cash_amount=100,mp_amount=20"))
    with pytest.raises(DataError), scratch_database.begin() as db:
        db.execute(text("""INSERT INTO salary_payments(entry_key,entry_date,employee,amount,version)
            VALUES ('SIM-BAD','2026-01-01','Otra',10,1)"""))
