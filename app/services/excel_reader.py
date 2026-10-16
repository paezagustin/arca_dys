"""Lectura acotada del formato Negocio. Nunca evalúa fórmulas ni modifica el libro."""
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
import re
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from app.services.business import money, record_date

MAX_UPLOAD = 5 * 1024 * 1024
HEADERS = ["FECHA", "VENTAS DEL DIA", "COMPRAS EN BLANCO", "MONTO", "COMPRAS EN NEGRO", "MONTO2"]


def digest(value):
    return sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def cell_record(cell, cached):
    return {"value": scalar(cell.value), "type": cell.data_type,
            "cached": scalar(cached.value) if cell.data_type == "f" else None}


def excel_money(cell, *, sale=False):
    if cell.data_type in {"f", "e", "b"} or isinstance(cell.value, bool):
        raise ValueError("Importe con fórmula, error o tipo no numérico: confirmar manualmente")
    if isinstance(cell.value, str) and not re.fullmatch(r"\d+(?:[.,]\d{1,2})?", cell.value.strip()):
        raise ValueError("Importe textual ambiguo: confirmar manualmente")
    return money(cell.value, allow_zero=sale)


def excel_date(cell):
    if isinstance(cell.value, datetime):
        return cell.value.date()
    if isinstance(cell.value, date):
        return cell.value
    return record_date(cell.value)


def read_excel(content, dataset):
    try:
        return _read_excel(content, dataset)
    except (ParseError, DefusedXmlException, OverflowError) as error:
        raise ValueError("El libro contiene XML o fechas inválidos; no se importaron datos") from error


def _read_excel(content, dataset):
    if not content or len(content) > MAX_UPLOAD:
        raise ValueError("El archivo debe ser XLSX y pesar hasta 5 MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            files = archive.infolist()
            if len(files) > 500 or sum(f.file_size for f in files) > 50 * 1024 * 1024 or any("vbaproject" in f.filename.lower() for f in files):
                raise ValueError("El libro excede los límites de lectura o contiene macros")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, KeyError, OSError, ValueError) as error:
        raise ValueError("No se pudo leer un XLSX válido, sin macros y dentro de los límites") from error
    try:
        names = {name.strip().casefold(): name for name in book.sheetnames}
        if "ingresos y egresos" not in names or "dashboard" not in names:
            raise ValueError("El libro debe contener las hojas ingresos y egresos y dashboard")
        detail = book[names["ingresos y egresos"]]
        if [str(detail.cell(1, column).value or "").strip().upper() for column in range(1, 7)] != HEADERS:
            raise ValueError("Los encabezados del detalle no coinciden con el formato esperado")
        items, totals = [], defaultdict(lambda: defaultdict(lambda: Decimal("0.00")))
        warnings = []

        def add(sheet, number, slot, kind, raw, proposed, state, issue=None):
            items.append({"sheet": sheet, "row_number": number, "slot": slot, "kind": kind, "raw": raw,
                "proposed": proposed, "state": state, "issue": issue,
                "source_key": digest([dataset, sheet, number, slot]), "fingerprint": digest(raw)})

        for sheet_key in ("ingresos y egresos", "dashboard"):
            sheet = book[names[sheet_key]]
            cache = cached[sheet.title]
            if sheet.max_row > 10000 or sheet.max_column > 30:
                raise ValueError("La hoja excede 10000 filas o 30 columnas")
            # Dos iteradores mantienen lectura lineal; cell() repetido en read_only sería cuadrático.
            for number, (row, cached_row) in enumerate(zip(sheet.iter_rows(), cache.iter_rows()), 1):
                raw = {cell.coordinate: cell_record(cell, cached_row[index]) for index, cell in enumerate(row) if cell.value is not None}
                if not raw:
                    continue
                add(sheet.title, number, "source", "source" if sheet_key == "ingresos y egresos" else "historical",
                    raw, {}, "source")
                if sheet_key == "dashboard" or number == 1:
                    continue
                for slot, label_column, amount_column, kind, classification in (("B", None, 1, "sale", ""),
                        ("D", 2, 3, "expense", "COMPRAS EN BLANCO"), ("F", 4, 5, "expense", "COMPRAS EN NEGRO")):
                    label = str(row[label_column].value or "").strip() if label_column is not None else ""
                    if row[amount_column].value is None and not label:
                        continue
                    proposed = {"kind": kind, "category": "Ventas" if kind == "sale" else "Compras",
                        "medium": "mixed" if kind == "sale" else "unspecified", "classification": classification,
                        "description": label, "label": label, "active": True}
                    issues = []
                    try:
                        proposed["entry_date"] = excel_date(row[0]).isoformat()
                    except (ValueError, TypeError):
                        issues.append("Fecha ausente o ambigua")
                    try:
                        proposed["amount"] = str(excel_money(row[amount_column], sale=kind == "sale"))
                    except (ValueError, TypeError, ArithmeticError) as error:
                        issues.append(str(error))
                    group = {key: value for key, value in raw.items() if key in {f"A{number}", f"{slot}{number}", f"{chr(65+label_column)}{number}" if label_column is not None else ""}}
                    add(sheet.title, number, slot, kind, group, proposed, "review" if issues else "ready", "; ".join(issues) or
                        ("Compra sin descripción: se conserva sin proveedor" if kind == "expense" and not label else None))
                    # Control independiente: suma de celdas numéricas válidas, incluso si falta una etiqueta.
                    if "amount" in proposed and "entry_date" in proposed:
                        key = "sales" if kind == "sale" else slot
                        totals[proposed["entry_date"][:7]][key] += Decimal(proposed["amount"])
        controls = []
        dash = book[names["dashboard"]]
        dash_cache = cached[dash.title]
        for row in dash.iter_rows(min_row=3, max_row=min(dash.max_row, 14), min_col=3, max_col=11):
            if not isinstance(row[0].value, (date, datetime)):
                continue
            month = row[0].value.strftime("%Y-%m")
            amounts = totals.get(month, {})
            expected = {"D": amounts.get("sales", Decimal(0)), "E": amounts.get("D", Decimal(0)), "F": amounts.get("F", Decimal(0))}
            expected["G"] = expected["E"] + expected["F"]
            expected["K"] = expected["D"] - expected["G"]
            for column, calculated in expected.items():
                coordinate = f"{column}{row[0].row}"
                source = dash[coordinate]
                value = dash_cache[coordinate].value if source.data_type == "f" else source.value
                try:
                    cache_amount = money(value, allow_zero=True) if value is not None and Decimal(str(value)) >= 0 else Decimal(str(value))
                    delta = str(cache_amount - calculated) if cache_amount.is_finite() else None
                except (ValueError, TypeError, ArithmeticError):
                    delta = None
                controls.append({"cell": coordinate, "month": month, "calculated": str(calculated), "cached": scalar(value), "difference": delta})
        warnings.append("Las fórmulas del dashboard se conservan como referencia y no se ejecutan. Las cantidades por rangos fijos pueden diferir del detalle por fecha.")
        warnings.append("Los bloques manuales RECIBIDO/EMITIDO se conservan separados; no se interpretan como ventas ni gastos y no se suman a gestión.")
        return items, {"schema": 1, "source_totals": {month: {key: str(value) for key, value in values.items()} for month, values in sorted(totals.items())},
            "dashboard_controls": controls, "warnings": warnings, "excluded_sheets": [name for name in book.sheetnames if name not in {detail.title, dash.title}]}
    finally:
        book.close()
        cached.close()
