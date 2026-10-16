"""Exportación del período con el formato de ingresos y egresos del libro original."""
from collections import defaultdict
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from app.models.business import BusinessEntry, Supplier

HEADERS = ("FECHA", "VENTAS DEL DIA", "COMPRAS EN BLANCO", "MONTO", "COMPRAS EN NEGRO", "MONTO2")


def export_business(db, start, end):
    entries = db.execute(select(BusinessEntry, Supplier.name).outerjoin(Supplier, Supplier.id == BusinessEntry.supplier_id).where(
        BusinessEntry.active.is_(True), BusinessEntry.entry_date.between(start, end)).order_by(BusinessEntry.entry_date, BusinessEntry.id))
    days = defaultdict(lambda: {"sale": None, "white": [], "black": []})
    for entry, supplier in entries:
        if entry.kind == "expense" and entry.expense_type not in {"white", "black"}:
            raise ValueError("Hay egresos anteriores sin tipo. Editalos en Gestión y elegí En Blanco o En Negro antes de exportar.")
        # Excel solo conserva 15 cifras significativas. Rechazar pérdida silenciosa.
        if abs(entry.amount * 100) >= Decimal("1000000000000000"):
            raise ValueError("Un importe supera la precisión de Excel; no se exportó para evitar pérdida de centavos.")
        if entry.kind == "sale":
            days[entry.entry_date]["sale"] = entry.amount
        else:
            label = " · ".join(value for value in (supplier, entry.description) if value) or "Sin proveedor"
            days[entry.entry_date][entry.expense_type].append((label, entry.amount))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ingresos y egresos"
    sheet.append(HEADERS)
    widths = (12, 19, 32, 19, 32, 19)
    for cell, width in zip(sheet[1], widths):
        cell.fill = PatternFill("solid", fgColor="ADC5E7")
        cell.font = Font(name="Calibri", size=11, bold=True, color="192C3D")
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[cell.column_letter].width = width
    day = start
    while day <= end:
        values = days[day]
        for n in range(max(len(values["white"]), len(values["black"]), 1)):
            white = values["white"][n] if n < len(values["white"]) else (None, None)
            black = values["black"][n] if n < len(values["black"]) else (None, None)
            sheet.append([day, values["sale"] if n == 0 else None, *white, *black])
            for cell in sheet[sheet.max_row]:
                cell.font = Font(name="Calibri", size=11, bold=cell.column in {3, 5}, color="192C3D")
                if cell.column == 1:
                    cell.number_format = "dd/mm/yyyy"
                elif cell.column in {2, 4, 6}:
                    cell.number_format = '"$ "#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif cell.value is not None:
                    # Conservar nombres/descripciones literales, nunca fórmulas del usuario.
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(cell.value))
                    cell.data_type = "s"
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[sheet.max_row].height = 30
        day += timedelta(days=1)
    table = Table(displayName="Tabla2", ref=f"A1:F{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = table.ref
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
